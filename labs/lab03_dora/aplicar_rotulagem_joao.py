"""Aplica o CSV auditável de João à planilha da amostra-ouro, sem editar células à mão.

    python -m labs.lab03_dora.aplicar_rotulagem_joao
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from labs.lab03_dora.amostra_ouro import CORRECTIVE, GOLD_DIR, PROJECT_TYPES, REAL_DELIVERY, read_labels

CSV_PATH = GOLD_DIR / "rotulagem_joao.csv"
XLSX_PATH = GOLD_DIR / "rotulagem_joao.xlsx"
REQUIRED_COLUMNS = {"escopo", "repositorio", "tag", "dimensao", "rotulo", "proveniencia"}
ALLOWED = {
    "tipo_projeto": set(PROJECT_TYPES),
    "entregas_reais": set(REAL_DELIVERY),
    "corretiva": set(CORRECTIVE),
}


def apply(csv_path: Path = CSV_PATH, xlsx_path: Path = XLSX_PATH) -> None:
    with csv_path.open(newline="", encoding="utf-8") as source:
        reader = csv.DictReader(source)
        if not REQUIRED_COLUMNS.issubset(reader.fieldnames or []):
            raise ValueError("CSV sem as colunas de chave, rótulo e proveniência")
        labels = {}
        for row in reader:
            key = (row["escopo"], row["repositorio"], row["tag"], row["dimensao"])
            if key in labels:
                raise ValueError(f"Chave duplicada no CSV: {key}")
            if row["rotulo"] not in ALLOWED.get(row["dimensao"], set()):
                raise ValueError(f"Rótulo inválido no CSV: {key}")
            if row["proveniencia"] != "humano_assistido_ia":
                raise ValueError(f"Proveniência inesperada no CSV: {key}")
            labels[key] = row["rotulo"]
    if len(labels) != 420:
        raise ValueError(f"Esperados 420 rótulos, encontrados {len(labels)}")

    workbook = load_workbook(xlsx_path)
    repositories = workbook["Repositorios"]
    releases = workbook["Releases"]
    if repositories.max_row != 61 or releases.max_row != 301:
        raise ValueError("A planilha não contém 60 repositórios e 300 releases")
    remaining = set(labels)
    for row in repositories.iter_rows(min_row=2):
        repository = row[0].value
        for column, dimension in ((4, "tipo_projeto"), (5, "entregas_reais")):
            key = ("repositorios", repository, "", dimension)
            if key not in remaining:
                raise ValueError(f"Chave não encontrada no CSV: {key}")
            row[column].value = labels[key]
            remaining.remove(key)
    for row in releases.iter_rows(min_row=2):
        key = ("releases", row[0].value, row[1].value, "corretiva")
        if key not in remaining:
            raise ValueError(f"Chave não encontrada no CSV: {key}")
        row[7].value = labels[key]
        remaining.remove(key)
    if remaining:
        raise ValueError(f"{len(remaining)} chaves do CSV sem linha na planilha")

    instructions = workbook["Instrucoes"]
    instructions["A1"] = "Rotulagem de João revisada por humano com assistência de IA; proveniência por linha em rotulagem_joao.csv."
    instructions["A2"] = "João aceitou as sugestões válidas e definiu como sim as 163 sugestões incertas de corretividade."
    temporary = xlsx_path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        read_labels(temporary)
        temporary.replace(xlsx_path)
    finally:
        temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    apply()
    print(f"OK: 420 rótulos aplicados em {XLSX_PATH}")
