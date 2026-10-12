"""Amostra-ouro (secao 6 do enunciado, issue #67): sorteio reprodutivel e planilhas de rotulagem.

Sorteia 60 repositorios da amostra final e 5 releases de cada um, com semente fixa, e gera uma planilha
por avaliador. As planilhas nao mostram a saida da heuristica de release corretiva.

    python -m labs.lab03_dora.amostra_ouro --config labs/lab03_dora/config.yaml
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Sequence
from urllib.parse import quote

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from labs.lab03_dora.config import ConfigError, load_config
from labs.lab03_dora.releases import compare_commits
from shared.github_client import load_env_file
from shared.github_rest import GitHubAPIError, GitHubNotFoundError, GitHubRestClient, ResponseCache

SEED = 42
REPOSITORIES_IN_SAMPLE = 60
RELEASES_PER_REPOSITORY = 5
EVALUATORS = ("pedro", "joao", "salomao")
PROJECT_TYPES = ("biblioteca/framework", "aplicacao/servico", "ferramenta CLI", "outro")
REAL_DELIVERY = ("sim", "nao", "incerto")
CORRECTIVE = ("sim", "nao")
LAB_DIR = Path(__file__).resolve().parent
GOLD_DIR = LAB_DIR / "amostra_ouro"

REPO_HEADERS = ("repositorio", "link", "descricao", "topicos", "tipo_projeto", "entregas_reais", "observacoes")
RELEASE_HEADERS = (
    "repositorio", "tag", "publicada_em", "link_release", "link_commits", "notas_da_release", "commits",
    "corretiva", "observacoes",
)
TEXT_COLUMNS = {"descricao": 50, "topicos": 30, "notas_da_release": 80, "commits": 70}
MAX_NOTES_CHARS = 4000
MAX_COMMITS_LISTED = 60
INSTRUCTIONS = (
    "Rotulagem independente da amostra-ouro (secao 6 do enunciado).",
    "Nao consulte os outros avaliadores nem a saida da heuristica de release corretiva.",
    "Aba Repositorios: tipo do projeto (biblioteca/framework, aplicacao/servico, ferramenta CLI ou outro) e se as "
    "releases representam entregas reais ao usuario (sim, nao ou incerto).",
    "Aba Releases: para cada release sorteada, se ela e corretiva (sim/nao), isto e, se o objetivo principal e "
    "corrigir um defeito recente. As colunas notas_da_release e commits trazem o texto da release e a primeira "
    "linha de cada commit desde a release anterior; os links abrem a pagina completa no GitHub.",
    "Use as listas de opcoes das colunas; observacoes sao livres. Salve o arquivo com o mesmo nome e faca o commit.",
)


def draw_repositories(repositories: pd.DataFrame, n: int = REPOSITORIES_IN_SAMPLE) -> pd.DataFrame:
    if len(repositories) < n:
        raise ValueError(f"A amostra final tem {len(repositories)} repositorios; o sorteio precisa de {n}.")
    return repositories.sample(n=n, random_state=SEED)[["repository"]].reset_index(drop=True)


def draw_releases(releases: pd.DataFrame, repositories: pd.Series, per_repository: int = RELEASES_PER_REPOSITORY) -> pd.DataFrame:
    """Ate 5 releases da janela por repositorio sorteado, na ordem do sorteio de repositorios."""
    chosen = []
    for repo in repositories:
        own = releases[releases["repository"] == repo].sort_values(["published_at", "tag_name"])
        if own.empty:
            raise ValueError(f"{repo} nao tem releases em releases.csv; rode a etapa releases.")
        picked = own.sample(n=min(per_repository, len(own)), random_state=SEED).sort_values("published_at")
        chosen.append(picked)
    return pd.concat(chosen, ignore_index=True)[["repository", "tag_name", "published_at", "previous_tag"]]


def release_link(repo: str, tag: str) -> str:
    return f"https://github.com/{repo}/releases/tag/{quote(tag, safe='')}"


def commits_link(repo: str, tag: str, previous: str | None) -> str:
    if not previous:
        return f"https://github.com/{repo}/commits/{quote(tag, safe='')}"
    return f"https://github.com/{repo}/compare/{quote(previous, safe='')}...{quote(tag, safe='')}"


def cell_text(value: str | None, limit: int = MAX_NOTES_CHARS) -> str:
    """Texto seguro para uma celula: sem caracteres de controle, sem virar formula e com tamanho limitado."""
    text = ILLEGAL_CHARACTERS_RE.sub("", (value or "").replace("\r\n", "\n")).strip()
    if len(text) > limit:
        text = text[:limit].rstrip() + "\n[... texto cortado; veja o link]"
    return " " + text if text.startswith("=") else text  # openpyxl grava texto com "=" como formula


def commit_summary(messages: list[str]) -> str:
    lines = [f"- {message.splitlines()[0] if message else ''}" for message in messages[:MAX_COMMITS_LISTED]]
    if len(messages) > MAX_COMMITS_LISTED:
        lines.append(f"(+{len(messages) - MAX_COMMITS_LISTED} commits; veja o link)")
    return cell_text("\n".join(lines))


def enrich(client, repositories: pd.DataFrame, releases: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Traz descricao/topicos dos repositorios e notas/commits das releases (respostas ficam no cache)."""
    repositories, releases = repositories.copy(), releases.copy()
    descriptions, topics = [], []
    for repo in repositories["repository"]:
        data = client.get(f"/repos/{repo}").data
        descriptions.append(cell_text(data.get("description")))
        topics.append(", ".join(data.get("topics") or []))
    repositories["description"], repositories["topics"] = descriptions, topics
    notes, commits = [], []
    for r in releases.itertuples():
        try:
            body = client.get(f"/repos/{r.repository}/releases/tags/{quote(r.tag_name, safe='')}").data.get("body")
        except GitHubNotFoundError:
            body = "(release nao encontrada na API)"
        notes.append(cell_text(body) or "(sem notas)")
        previous = r.previous_tag if isinstance(r.previous_tag, str) and r.previous_tag else None
        if previous is None:
            commits.append("(primeira release: sem release anterior para comparar)")
            continue
        try:
            found, _ = compare_commits(client, r.repository, previous, r.tag_name)
            commits.append(commit_summary([c.get("commit", {}).get("message", "") for c in found]) or "(sem commits)")
        except (GitHubNotFoundError, RuntimeError):
            commits.append("(compare indisponivel; use o link)")
    releases["notes"], releases["commits"] = notes, commits
    return repositories, releases


def _sheet(workbook: Workbook, title: str, headers: Sequence[str], rows: list[list], choices: dict[str, Sequence[str]]):
    sheet = workbook.create_sheet(title)
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        if header in TEXT_COLUMNS:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = max(len(rows) + 1, 2)
    for column, options in choices.items():
        letter = sheet.cell(row=1, column=headers.index(column) + 1).column_letter
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                    showErrorMessage=True, error="Escolha um valor da lista.")
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last}")
        for cells in sheet[f"{letter}2:{letter}{last}"]:
            cells[0].fill = PatternFill("solid", fgColor="FFF2CC")
    for index, header in enumerate(headers, start=1):
        width = TEXT_COLUMNS.get(header) or (60 if header.startswith("link") else max(14, len(header) + 4))
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"
    return sheet


def write_sheet(path: Path, repositories: pd.DataFrame, releases: pd.DataFrame) -> Path:
    workbook = Workbook()
    intro = workbook.active
    intro.title = "Instrucoes"
    for line in INSTRUCTIONS:
        intro.append([line])
    intro.column_dimensions["A"].width = 140
    repo_rows = [[r.repository, f"https://github.com/{r.repository}", getattr(r, "description", ""),
                  getattr(r, "topics", ""), None, None, None] for r in repositories.itertuples()]
    _sheet(workbook, "Repositorios", REPO_HEADERS, repo_rows,
           {"tipo_projeto": PROJECT_TYPES, "entregas_reais": REAL_DELIVERY})
    release_rows = [
        [r.repository, r.tag_name, r.published_at, release_link(r.repository, r.tag_name),
         commits_link(r.repository, r.tag_name, r.previous_tag if isinstance(r.previous_tag, str) else None),
         getattr(r, "notes", ""), getattr(r, "commits", ""), None, None]
        for r in releases.itertuples()
    ]
    _sheet(workbook, "Releases", RELEASE_HEADERS, release_rows, {"corretiva": CORRECTIVE})
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def read_labels(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Le uma planilha preenchida e valida os rotulos (usado no calculo do kappa, issue #71)."""
    workbook = load_workbook(path, read_only=True)
    frames = []
    for title, headers, checks in (
        ("Repositorios", REPO_HEADERS, {"tipo_projeto": PROJECT_TYPES, "entregas_reais": REAL_DELIVERY}),
        ("Releases", RELEASE_HEADERS, {"corretiva": CORRECTIVE}),
    ):
        rows = list(workbook[title].iter_rows(values_only=True))
        frame = pd.DataFrame(rows[1:], columns=rows[0])
        if tuple(frame.columns) != headers:
            raise ValueError(f"{path.name}: colunas inesperadas na aba {title}.")
        for column, options in checks.items():
            invalid = frame[~frame[column].isin(options)]
            if not invalid.empty:
                raise ValueError(f"{path.name}: {len(invalid)} linha(s) da aba {title} sem rotulo valido em {column} "
                                 f"(ex.: linha {invalid.index[0] + 2}).")
        frames.append(frame)
    workbook.close()
    return frames[0], frames[1]


def generate(output_dir: Path, gold_dir: Path, evaluators: Sequence[str] = EVALUATORS, force: bool = False,
             client=None) -> list[Path]:
    repositories = draw_repositories(pd.read_csv(output_dir / "repositories.csv"))
    releases = draw_releases(pd.read_csv(output_dir / "releases.csv", dtype={"tag_name": str, "previous_tag": str}),
                             repositories["repository"])
    targets = [gold_dir / f"rotulagem_{name}.xlsx" for name in evaluators]
    existing = [path.name for path in targets if path.exists()]
    if existing and not force:
        raise FileExistsError(f"Planilhas ja existem ({', '.join(existing)}); use --forcar para sobrescrever.")
    gold_dir.mkdir(parents=True, exist_ok=True)
    repositories.to_csv(gold_dir / "repositorios_sorteados.csv", index=False)
    releases.to_csv(gold_dir / "releases_sorteadas.csv", index=False)
    if client is not None:
        repositories, releases = enrich(client, repositories, releases)
    return [write_sheet(path, repositories, releases) for path in targets]


def default_client(config) -> GitHubRestClient | None:
    load_env_file(LAB_DIR.parents[1] / ".env")
    token = os.environ.get("GITHUB_TOKEN")
    if not token:
        return None
    return GitHubRestClient(token, cache=ResponseCache(config.cache_dir), max_retries=config.max_retries,
                            timeout=config.timeout)


def main(argv: Sequence[str] | None = None, client_factory=default_client) -> int:
    parser = argparse.ArgumentParser(description="Sorteia a amostra-ouro e gera as planilhas de rotulagem.")
    parser.add_argument("--config", type=Path, default=LAB_DIR / "config.yaml")
    parser.add_argument("--saida", type=Path, default=GOLD_DIR)
    parser.add_argument("--avaliadores", nargs="+", default=list(EVALUATORS))
    parser.add_argument("--forcar", action="store_true", help="sobrescreve planilhas existentes (perde rotulos!)")
    args = parser.parse_args(argv)
    try:
        config = load_config(args.config)
        client = client_factory(config)
        if client is None:
            print("Sem GITHUB_TOKEN: planilhas sem notas e commits (so com os links).")
        paths = generate(config.output_dir, args.saida, args.avaliadores, args.forcar, client)
    except (ConfigError, FileNotFoundError, FileExistsError, ValueError, GitHubAPIError) as error:
        print(f"Nao foi possivel gerar a amostra-ouro: {error}")
        return 1
    for path in paths:
        print(f"OK: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
