from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from labs.lab03_dora import amostra_ouro


def write_inputs(folder, repos=70, releases_per_repo=7):
    names = [f"org/repo{i:02d}" for i in range(repos)]
    pd.DataFrame({"repository": names, "default_branch": "main"}).to_csv(folder / "repositories.csv", index=False)
    rows = []
    for name in names:
        for j in range(releases_per_repo if name != "org/repo00" else 2):
            rows.append({"repository": name, "tag_name": f"v1.{j}", "published_at": f"2026-0{1 + j % 9}-10T00:00:00Z",
                         "previous_tag": f"v1.{j - 1}" if j else ""})
    pd.DataFrame(rows).to_csv(folder / "releases.csv", index=False)


def test_draw_is_reproducible_with_the_documented_seed(tmp_path):
    write_inputs(tmp_path)
    frame = pd.read_csv(tmp_path / "repositories.csv")
    first = amostra_ouro.draw_repositories(frame)
    assert len(first) == 60 and first["repository"].is_unique
    assert first.equals(amostra_ouro.draw_repositories(frame))
    assert list(first["repository"]) == list(frame.sample(n=60, random_state=42)["repository"])
    with pytest.raises(ValueError, match="60"):
        amostra_ouro.draw_repositories(frame.head(10))


def test_five_releases_per_repository_or_all_when_fewer(tmp_path):
    write_inputs(tmp_path)
    releases = pd.read_csv(tmp_path / "releases.csv", dtype={"tag_name": str, "previous_tag": str})
    drawn = amostra_ouro.draw_releases(releases, pd.Series(["org/repo00", "org/repo01"]))
    assert drawn.groupby("repository").size().to_dict() == {"org/repo00": 2, "org/repo01": 5}
    with pytest.raises(ValueError, match="sem|nao tem"):
        amostra_ouro.draw_releases(releases, pd.Series(["org/inexistente"]))


def test_links_encode_tags_and_fall_back_to_commit_history():
    assert amostra_ouro.release_link("o/r", "release/1.0") == "https://github.com/o/r/releases/tag/release%2F1.0"
    assert amostra_ouro.commits_link("o/r", "v2", "v1") == "https://github.com/o/r/compare/v1...v2"
    assert amostra_ouro.commits_link("o/r", "v1", None) == "https://github.com/o/r/commits/v1"


def test_generates_one_sheet_per_evaluator_with_dropdowns_and_no_heuristic(tmp_path):
    write_inputs(tmp_path)
    gold = tmp_path / "ouro"
    paths = amostra_ouro.generate(tmp_path, gold)

    assert [p.name for p in paths] == ["rotulagem_pedro.xlsx", "rotulagem_joao.xlsx", "rotulagem_salomao.xlsx"]
    assert (gold / "repositorios_sorteados.csv").is_file() and (gold / "releases_sorteadas.csv").is_file()
    book = load_workbook(paths[0])
    assert book.sheetnames == ["Instrucoes", "Repositorios", "Releases"]
    repos, releases = book["Repositorios"], book["Releases"]
    assert repos.max_row == 61 and releases.max_row > 61
    assert releases["D2"].value.startswith("https://github.com/org/repo")
    formulas = {v.formula1 for v in releases.data_validations.dataValidation}
    assert '"sim,nao"' in formulas
    headers = [c.value for c in releases[1]]
    assert not any("heuristica" in str(h) for h in headers)

    with pytest.raises(FileExistsError):
        amostra_ouro.generate(tmp_path, gold)


def test_read_labels_accepts_complete_sheets_and_rejects_blanks(tmp_path):
    write_inputs(tmp_path)
    path = amostra_ouro.generate(tmp_path, tmp_path / "ouro", ["pedro"])[0]
    with pytest.raises(ValueError, match="tipo_projeto"):
        amostra_ouro.read_labels(path)

    book = load_workbook(path)
    for row in range(2, book["Repositorios"].max_row + 1):
        book["Repositorios"].cell(row=row, column=3, value="biblioteca/framework")
        book["Repositorios"].cell(row=row, column=4, value="sim")
    for row in range(2, book["Releases"].max_row + 1):
        book["Releases"].cell(row=row, column=6, value="nao")
    book.save(path)
    repos, releases = amostra_ouro.read_labels(path)
    assert len(repos) == 60 and set(releases["corretiva"]) == {"nao"}


def test_command_line_reports_missing_inputs(tmp_path):
    config = tmp_path / "config.yaml"
    config.write_text("janela:\n  inicio: 2025-10-01\n  fim: 2026-09-30\nsaida:\n  dados: .\n", encoding="utf-8")
    assert amostra_ouro.main(["--config", str(config), "--saida", str(tmp_path / "ouro")]) == 1
    write_inputs(tmp_path)
    assert amostra_ouro.main(["--config", str(config), "--saida", str(tmp_path / "ouro"), "--avaliadores", "x"]) == 0
